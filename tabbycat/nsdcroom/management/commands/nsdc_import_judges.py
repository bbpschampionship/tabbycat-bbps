"""Fill the NSDC judge slots from the judge-allocation workbook.

Sheet 1: Batch Code, ..., Judge, Judge Contact  -> names that batch's judge slot
(keeps its private URL). "Backup Judges" sheet: No., Judge, Contact -> standby
judges with private URLs and no batch, ready to be moved onto any batch.
Contacts are stored on the judge (Person.phone) for the tab team only; the public
pages never show them. Safe to re-run.
"""
import re

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from participants.models import Adjudicator, Institution
from privateurls.utils import populate_url_keys
from tournaments.models import Tournament


def tidy(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


class Command(BaseCommand):
    help = "Import NSDC judge allocation (batch judges + backups)"

    def add_arguments(self, parser):
        parser.add_argument('workbook')
        parser.add_argument('--tournament', required=True)
        parser.add_argument('--dry-run', action='store_true')

    def handle(self, *args, **opts):
        import openpyxl
        t = Tournament.objects.get(slug=opts['tournament'])
        wb = openpyxl.load_workbook(opts['workbook'], read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        head = [tidy(h).lower() for h in rows[0]]
        try:
            ci, ji, pi = head.index('batch code'), head.index('judge'), head.index('judge contact')
        except ValueError:
            raise CommandError(f"Unexpected header: {rows[0]}")
        alloc = {}
        for r in rows[1:]:
            if r[ci] and tidy(r[ji]):
                alloc[str(r[ci]).strip().zfill(2)] = (tidy(r[ji]), tidy(r[pi]))
        backups = []
        for ws in wb.worksheets[1:]:
            if 'backup' in ws.title.lower():
                for r in list(ws.iter_rows(values_only=True))[1:]:
                    if r and len(r) > 2 and tidy(r[1]):
                        backups.append((tidy(r[1]), tidy(r[2])))

        rooms = {i.code: i for i in Institution.objects.filter(team__tournament=t).distinct()}
        missing = sorted(set(alloc) - set(rooms))
        if missing:
            raise CommandError(f"Batch codes not in tournament: {missing}")
        self.stdout.write(f"{len(alloc)} batch judges, {len(backups)} backups")
        if opts['dry_run']:
            for code, (name, _) in sorted(alloc.items()):
                self.stdout.write(f"  {code} {rooms[code].name}: {name}")
            for name, _ in backups:
                self.stdout.write(f"  backup: {name}")
            return

        with transaction.atomic():
            for code, (name, phone) in sorted(alloc.items()):
                judge = Adjudicator.objects.filter(tournament=t, institution=rooms[code]).order_by('id').first()
                if judge is None:
                    judge = Adjudicator(tournament=t, institution=rooms[code])
                judge.name, judge.phone = name, phone
                judge.save()
            for name, phone in backups:
                judge = Adjudicator.objects.filter(tournament=t, institution__isnull=True, name=name).first() \
                    or Adjudicator(tournament=t, institution=None, name=name)
                judge.phone = phone
                judge.save()
            populate_url_keys(Adjudicator.objects.filter(tournament=t, url_key__isnull=True))
        self.stdout.write(self.style.SUCCESS(
            f"Judges set: {Adjudicator.objects.filter(tournament=t, institution__isnull=False).exclude(name='').count()} "
            f"batch judges named, {Adjudicator.objects.filter(tournament=t, institution__isnull=True).count()} backups"))
